"""Offline source contracts; every item here is invented test data."""
import sys
import unittest
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

class SourcesTest(unittest.TestCase):
    def test_sources_are_disabled_and_manual_source_extends_registry(self):
        from src.sources import default_registry, ManualJSONSource
        registry = default_registry()
        self.assertEqual(registry.search('bangumi', '示例作品')['status'], 'disabled')
        registry.register(ManualJSONSource('custom', [{'source_id': 'demo-01', 'name': '示例作品', 'summary': '自填资料'}]))
        self.assertEqual(registry.search('custom', '示例')['candidates'][0]['source_id'], 'demo-01')
        self.assertEqual(registry.detail('custom', 'demo-01')['info']['summary'], '自填资料')

    def test_bangumi_detail_and_search_produce_generic_info(self):
        from src.sources import default_registry
        def transport(url, payload=None, headers=None):
            subject = {'id': 101, 'name': '示例作品', 'name_cn': '示例作品', 'type': 2, 'summary': '示例简介', 'images': {'large': 'https://example.invalid/cover.png'}}
            return {'data': [subject], 'total': 1} if payload else subject
        registry = default_registry({'bangumi': {'enabled': True, 'min_interval': 0}}, {'bangumi': transport})
        candidate = registry.search('Bangumi', '示例作品')['candidates'][0]
        self.assertEqual(candidate['source_id'], '101')
        info = registry.detail('bangumi', '101')['info']
        self.assertEqual(info['type'], 'animation')
        self.assertFalse(info['cover_redistribution_allowed'])

    def test_manual_id_never_falls_back_and_existing_values_survive_failure(self):
        from src.sources import default_registry, enrich_record, ManualJSONSource
        registry = default_registry()
        registry.register(ManualJSONSource('custom', [{'source_id': 'a', 'name': '示例作品'}]))
        record = {'name': '人工名称', 'manual_fields': ['name'], 'source_id': 'missing', 'source': 'custom', 'source_id_origin': 'manual', 'web_info': {'summary': '保留简介'}}
        enrich_record(record, registry, sources=['custom'])
        self.assertEqual(record['web_info']['summary'], '保留简介')
        self.assertEqual(record['source_checks']['custom']['status'], 'genuine_not_found')
        record['source_id'] = 'a'
        enrich_record(record, registry, sources=['custom'], refresh=True)
        self.assertEqual(record['name'], '人工名称')

    def test_cooldown_remembers_status_without_repeating_request(self):
        from src.sources.base import Source, SourceRegistry
        from urllib.error import HTTPError
        class Busy(Source):
            key = 'busy'
            def _search(self, title):
                raise HTTPError('https://example.invalid', 429, 'busy', {'Retry-After': '120'}, None)
        registry = SourceRegistry([Busy(enabled=True)])
        self.assertEqual(registry.search('busy','x')['status'], 'rate_limited')
        self.assertEqual(registry.search('busy','x')['status'], 'cooldown')

    def test_quark_total_pagination_manual_override_and_ambiguity(self):
        from src.sources.quark import crawl_share, match_record_quark_folder, scan_records
        def page_loader(fid, page, size):
            if fid != '0': return {'data': {'list': []}}
            if page == 1:
                return {'metadata': {'_total': 3}, 'data': {'list': [{'fid':'a', 'file_name':'DEMO101', 'file_type':0}, {'fid':'b','file_name':'DEMO101','file_type':0}]}}
            return {'metadata': {'_total': 3}, 'data': {'list': [{'fid':'c','file_name':'DEMO102','file_type':0}]}}
        folders, requests = crawl_share(page_loader, max_depth=1, page_size=2)
        self.assertEqual(len(folders), 3)
        self.assertEqual(match_record_quark_folder({'code':'DEMO101'}, folders, '', '')[0], 'ambiguous')
        manual={'code':'DEMO101','quark_location_manual':True,'quark_location_url':'https://example.invalid/chosen'}
        self.assertEqual(match_record_quark_folder(manual, folders, '', '')[1], manual['quark_location_url'])
        blocked = scan_records([{'code':'DEMO102'}], page_loader=page_loader)
        self.assertEqual(blocked['status'], 'authorization_required')
        imported = scan_records([{'code':'DEMO102'}], listing=folders)
        self.assertEqual(imported['records'][0]['quark_location_status'], 'matched')

    def test_quark_network_failure_is_not_not_found(self):
        from src.sources.quark import scan_records
        def fail(*args): raise OSError('offline')
        result = scan_records([{'code':'DEMO103'}], page_loader=fail, authorized=True)
        self.assertEqual(result['records'][0]['quark_location_status'], 'network_error')

    def test_steam_official_catalog_uses_pagination_and_no_store_scraping(self):
        from src.sources import default_registry
        def transport(url, payload=None, headers=None):
            return {'response': {'apps': [{'appid': 701, 'name': '示例目录作品'}], 'have_more_results': False}}
        registry = default_registry({'steam': {'enabled':True, 'api_key':'test-placeholder', 'min_interval':0}}, {'steam':transport})
        self.assertEqual(registry.search('steam','示例目录作品')['candidates'][0]['source_id'], '701')
        result = registry.detail('steam','701')
        self.assertEqual(result['info']['url'], 'https://store.steampowered.com/app/701/')
        self.assertNotIn('cover', result['info'])

class HelperSourcesTest(unittest.TestCase):
    def test_helper_manual_json_enrich_and_failed_daily_step_preserve_other_results(self):
        from src.helper import server
        job=server.Job('offline-test')
        config={'custom_sources':[{'key':'custom','enabled':True,'records':[{'source_id':'demo-1','name':'示例手册','summary':'示例简介'}]}]}
        record={'id':'item-1','name':'示例手册','source':'custom','source_id':'demo-1','source_id_origin':'manual'}
        response=server.job_enrich_ids(job,{'records':[record],'sourceConfig':config})
        self.assertEqual(response['changed']['item-1']['web_info']['summary'],'示例简介')
        daily=server.job_daily_routine(job,{'records':[record],'sourceConfig':config,'sources':['custom'],'shares':[],'sites':[]})
        self.assertEqual(daily['enrich']['changed']['item-1']['web_info']['summary'],'示例简介')
        self.assertEqual(daily['refCovers']['status'],'unconfigured')
        failed=server.job_enrich_ids(job,{'records':[dict(record,source_id='missing',web_info={'summary':'旧简介'})],'sourceConfig':config})
        self.assertEqual(failed['changed']['item-1']['web_info']['summary'],'旧简介')

    def test_url_guard_blocks_private_and_redirect_targets(self):
        from src.sources.http import validate_public_url, PublicRedirectHandler
        for url in ('http://127.0.0.1/','http://[::1]/','http://192.168.1.2/','http://example.invalid:8080/'):
            with self.assertRaises(ValueError):validate_public_url(url)
        redirect=PublicRedirectHandler()
        from urllib.request import Request
        with self.assertRaises(ValueError):redirect.redirect_request(Request('https://example.invalid'),None,302,'',{},'http://127.0.0.1/')

    def test_explicit_disabled_sources_stay_disabled(self):
        from src.helper import server
        from src.sources import default_registry
        registry=server.registry_for_payload({'sources':['bangumi'],'sourceConfig':{'bangumi':{'enabled':False}}})
        self.assertFalse(registry.sources['bangumi'].enabled)
        custom=default_registry({'custom_sources':[{'key':'library','enabled':False,'records':[{'source_id':'demo-1','name':'示例'}]}]})
        self.assertEqual(custom.search('library','示例')['status'],'disabled')

    def test_daily_tally_matches_ui_and_partial_steps_keep_results(self):
        from src.helper import server
        from unittest.mock import patch
        config={'custom_sources':[{'key':'library','enabled':True,'records':[{'source_id':'demo-1','name':'示例资料','url':'https://example.invalid/item'}]}]}
        payload={'records':[{'id':'demo-a','name':'示例资料'}],'sourceConfig':config,'sources':['library'],'shares':[],
                 'sites':[{'id':'test-site','url':'https://example.invalid/','enabled':True}]}
        with patch('ref_sites.find_covers',return_value={'candidates':[{'id':'demo-a','image':'synthetic'}],'stats':{'matched':1,'networkFailed':1}}):
            response=server.job_daily_routine(server.Job('test'),payload)
        self.assertEqual(response['tally']['library']['new_matched'],1)
        self.assertEqual(response['status'],'partial')
        self.assertEqual(response['refCovers']['candidates'][0]['id'],'demo-a')
        self.assertEqual(payload['records'][0],{'id':'demo-a','name':'示例资料'})

    def test_helper_workbook_keeps_metadata_and_relation_tables(self):
        from src.helper import server
        from openpyxl import Workbook
        import io,base64,json
        workbook=Workbook();sheet=workbook.active;sheet.title='作品'
        sheet.append(['ID','名称','资料编号','资料来源','资料JSON','别名JSON'])
        sheet.append(['demo-a','示例资料','demo-ref','library',json.dumps({'summary':'示例简介'}),json.dumps(['别名'])])
        sources=workbook.create_sheet('来源配置');sources.append(['配置名','JSON']);sources.append(['sourceConfig','{"custom_sources":[]}'])
        folders=workbook.create_sheet('目录定位');folders.append(['合集ID','文件夹JSON']);folders.append(['collection-demo','[]'])
        output=io.BytesIO();workbook.save(output)
        response=server.job_legacy_import(server.Job('test'),{'fileName':'test.xlsx','fileBase64':base64.b64encode(output.getvalue()).decode(),'sources':[]})
        self.assertEqual(response['views'][0]['web_info']['summary'],'示例简介')
        self.assertEqual(response['views'][0]['aliases'],['别名'])
        self.assertIn('来源配置',response['workbook'])
        self.assertIn('目录定位',response['workbook'])

    def test_generic_csv_mapping_keeps_stable_ids_and_user_fields(self):
        from src.importers.column_mapping import import_table
        data='目录ID,名称,资料来源,资料编号,自定义备注\nitem-demo,示例资料,custom,demo-1,说明文字\n'
        meta,records=import_table(data.encode(),filename='sample.csv',mapping={'notes':'自定义备注'})
        self.assertEqual(records[0]['id'],'item-demo')
        self.assertEqual(records[0]['source_id'],'demo-1')
        self.assertEqual(records[0]['notes'],'说明文字')

if __name__ == '__main__': unittest.main()
