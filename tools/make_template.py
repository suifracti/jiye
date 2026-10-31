#!/usr/bin/env python3
"""Build the public workbook template with the same columns as the browser export."""
from pathlib import Path
import argparse
import re
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))

def make_template(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    script=(ROOT/'src/master/js/70-xlsx.js').read_text(encoding='utf-8')
    columns=script.split('const WORK_COLUMNS = [',1)[1].split('];',1)[0]
    headers=['ID']+[label for label,key in re.findall(r"\['([^']+)'\s*,\s*'([^']+)'\]",columns)]
    workbook=Workbook();works=workbook.active;works.title='作品';works.append(headers)
    # Empty template: users provide their own works, source configuration and links.
    sheets={
        '合集':[['ID','合集名称','下载链接1','下载链接2','作品数','合集完整JSON']],
        'HTML设置':[['配置项','设置值']],
        '说明':[['说明'],['先在作品表填写名称，已有作品请保留 ID；新增作品 ID 可留空。'],['资料来源、资料编号可填写自定义来源及编号；联网需要在母版中主动启用。'],['名称、版本、版本类型、附加标记和备注分列；JSON 列可留空。'],['标签用空格分隔；多项 JSON 字段使用 JSON 数组，例如 ["示例标签"]。'],['作品、合集、设置和分区可在母版确认后一次应用，并可撤销。'],['来源配置与项目配置仅填写自己有权提供的资料；不要随表格泄漏凭据。'],['CSV 仅包含作品表；完整往返使用 XLSX。']],
        '来源配置':[['配置名','JSON']],
        '目录定位':[['合集ID','文件夹JSON']],
        '分区':[['目录ID','分区名']],
        '项目配置':[['配置名','JSON']],
    }
    for name,rows in sheets.items():
        sheet=workbook.create_sheet(name)
        for row in rows:sheet.append(row)
    for sheet in workbook:
        sheet.freeze_panes='A2';sheet.sheet_view.showGridLines=False
        for cell in sheet[1]:
            cell.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='325A62');cell.alignment=Alignment(vertical='center',wrap_text=True)
        sheet.row_dimensions[1].height=30
        for row in sheet.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True);cell.number_format='@'
        for column in sheet.columns:
            label=str(column[0].value or '')
            sheet.column_dimensions[column[0].column_letter].width=65 if 'JSON' in label or label in ('说明','设置值','下载链接1','下载链接2') else 32 if label in ('名称','备注') else 19
        if sheet.title!='说明':sheet.auto_filter.ref=sheet.dimensions
    works.sheet_view.zoomScale=85
    validation=DataValidation(type='list',formula1='"是,否"',allow_blank=True)
    works.add_data_validation(validation)
    from openpyxl.utils import get_column_letter
    update_column=headers.index('更新')+1
    validation.add(f'{get_column_letter(update_column)}2:{get_column_letter(update_column)}2001')
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);workbook.save(path)
    return path

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,default=ROOT/'docs/表格模板.xlsx')
    args=parser.parse_args();print(make_template(args.output))
