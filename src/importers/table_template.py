"""Create a neutral workbook; no catalog values are shipped by this module."""
from .column_mapping import HEADERS

def make_template(path):
    from openpyxl import Workbook
    workbook=Workbook();works=workbook.active;works.title='作品'
    works.append([labels[0] for labels in HEADERS.values()])
    collections=workbook.create_sheet('合集');collections.append(['合集ID','名称','下载链接2','下载链接1','说明'])
    settings=workbook.create_sheet('HTML设置');settings.append(['设置名','值'])
    help_sheet=workbook.create_sheet('说明');help_sheet.append(['保留目录ID可与母版匹配；自定义列可在导入时映射。'])
    for sheet in workbook:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=22
    workbook.save(path)
    return path
